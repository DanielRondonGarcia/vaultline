"""Authenticated download renderers for the backup audit report.

These functions turn the already-sanitized report dictionary produced by
``build_backup_audit_report`` into downloadable bytes for JSON, CSV, Excel
(.xlsx), and PDF. They never touch persisted data, secrets, the filesystem, or
the job history; they only serialize the report they are given, so the T2
no-secrets invariant is preserved by construction.

Backup size is surfaced as two distinct per-metric aggregates captured from
the Restic JSON summary emitted at the end of each ``backup.run`` job:
``bytes_processed`` (source ``total_bytes_processed``) and ``bytes_added``
(source ``data_added``). Each metric is summed across in-period jobs that
reported an exact nonnegative integer value, and is ``None`` (never zero) when
no job reported it; per-metric ``reported_jobs``/``unavailable_jobs`` coverage
makes partial totals explicit. Metrics stay unavailable independently for
non-Restic strategies or jobs without an exact Restic summary.

The storage location is derived only from safe, non-credential-bearing
fields: ``BackupTargetRecord.path_storage`` (a relative POSIX sub-path) and the
linked ``StorageProfileRecord.backend_type``/``name``. Repository URLs, secret
IDs, runtime environment, and secret values are never exported.
"""

from __future__ import annotations

import csv
import io
import json
from typing import Any, Callable, Dict, Tuple

BACKUP_AUDIT_MATRIX_COLUMNS: Tuple[str, ...] = (
    "target_id",
    "service",
    "server_name",
    "server_host",
    "runtime_type",
    "backup_strategy",
    "backup_mode",
    "enabled",
    "schedule_cron",
    "schedule_source",
    "schedule_timezone",
    "retention_name",
    "keep_last",
    "keep_hourly",
    "keep_daily",
    "keep_weekly",
    "keep_monthly",
    "keep_yearly",
    "prune",
    "total",
    "completed",
    "succeeded",
    "failed",
    "other",
    "succeeded_percent",
    "failed_percent",
    "backup_bytes_processed",
    "backup_bytes_added",
    "backup_bytes_processed_reported_jobs",
    "backup_bytes_processed_unavailable_jobs",
    "backup_bytes_added_reported_jobs",
    "backup_bytes_added_unavailable_jobs",
    "storage_backend_type",
    "storage_profile_name",
    "path_storage",
)

BACKUP_AUDIT_XLSX_HEADERS: Tuple[str, ...] = (
    "ID del destino",
    "Servicio",
    "Servidor (worker)",
    "Host del servidor",
    "Tipo de ejecución",
    "Estrategia de respaldo",
    "Modo de respaldo",
    "Habilitado",
    "Programación (cron)",
    "Origen de programación",
    "Zona horaria",
    "Política de retención",
    "Últimos",
    "Cada hora",
    "Cada día",
    "Cada semana",
    "Cada mes",
    "Cada año",
    "Depuración",
    "Trabajos totales",
    "Completados",
    "Correctos",
    "Fallidos",
    "Otros",
    "Porcentaje de éxito",
    "Porcentaje de fallos",
    "Bytes de origen procesados",
    "Bytes agregados al backend",
    "Trabajos con bytes procesados",
    "Trabajos sin bytes procesados",
    "Trabajos con bytes agregados",
    "Trabajos sin bytes agregados",
    "Backend de almacenamiento",
    "Perfil de almacenamiento",
    "Ruta de almacenamiento",
)

_CONTENT_TYPES: Dict[str, str] = {
    "json": "application/json",
    "csv": "text/csv; charset=utf-8",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "pdf": "application/pdf",
}


def _spreadsheet_safe_cell(value: Any) -> Any:
    """Prevent untrusted text cells from being evaluated as spreadsheet formulas."""
    if not isinstance(value, str):
        return value
    if value.lstrip().startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def _flatten_target(row: Dict[str, Any]) -> list:
    retention = row.get("retention") or {}
    schedule = row.get("schedule") or {}
    server = row.get("server") or {}
    counts = row.get("job_counts") or {}
    location = row.get("storage_location") or {}
    backup_size = row.get("backup_size") or {}
    coverage = backup_size.get("coverage") or {}
    processed_coverage = coverage.get("bytes_processed") or {}
    added_coverage = coverage.get("bytes_added") or {}
    return [
        row.get("target_id"),
        row.get("service"),
        server.get("name"),
        server.get("host"),
        row.get("runtime_type"),
        row.get("backup_strategy"),
        row.get("backup_mode"),
        row.get("enabled"),
        schedule.get("cron_expression"),
        schedule.get("source"),
        schedule.get("timezone"),
        retention.get("name"),
        retention.get("keep_last"),
        retention.get("keep_hourly"),
        retention.get("keep_daily"),
        retention.get("keep_weekly"),
        retention.get("keep_monthly"),
        retention.get("keep_yearly"),
        retention.get("prune"),
        counts.get("total"),
        counts.get("completed"),
        counts.get("succeeded"),
        counts.get("failed"),
        counts.get("other"),
        row.get("succeeded_percent"),
        row.get("failed_percent"),
        backup_size.get("bytes_processed"),
        backup_size.get("bytes_added"),
        processed_coverage.get("reported_jobs"),
        processed_coverage.get("unavailable_jobs"),
        added_coverage.get("reported_jobs"),
        added_coverage.get("unavailable_jobs"),
        location.get("backend_type"),
        location.get("storage_profile_name"),
        location.get("path_storage"),
    ]


def _spanish_size_basis(available: bool) -> str:
    if not available:
        return "No hay datos históricos exactos de tamaño; los valores faltantes se marcan como no disponibles y no se estiman."
    return (
        "Los bytes de origen procesados (total_bytes_processed) y los bytes agregados al backend (data_added) "
        "provienen del resumen JSON de Restic de cada trabajo backup.run. Son métricas independientes; "
        "los valores faltantes se marcan como no disponibles, sin estimarlos."
    )


def _spanish_location_basis(available: bool) -> str:
    if not available:
        return "La ubicación de almacenamiento no está disponible."
    return "La ubicación combina la ruta del destino con el tipo y el nombre del perfil de almacenamiento vinculado."


def _spanish_coverage_status(status: str) -> str:
    return {"complete": "Completo", "partial": "Parcial", "unavailable": "No disponible"}.get(status, status)


def _spanish_job_status(status: Any) -> str:
    normalized = str(status or "").strip().lower()
    return {
        "succeeded": "correctos",
        "success": "correctos",
        "failed": "fallidos",
        "failure": "fallidos",
        "canceled": "cancelados",
        "cancelled": "cancelados",
        "pending": "pendientes",
        "queued": "en cola",
        "running": "en ejecución",
        "started": "iniciados",
    }.get(normalized, str(status or "desconocido"))


def _spanish_schedule_source(source: Any) -> str:
    normalized = str(source or "").strip().lower()
    return {
        "target": "destino",
        "settings": "configuración",
        "worker": "servidor",
        "global": "global",
        "default": "predeterminada",
    }.get(normalized, str(source or "no disponible"))


def export_backup_audit_json(report: Dict[str, Any]) -> bytes:
    return json.dumps(report, ensure_ascii=False, indent=2).encode("utf-8")


def export_backup_audit_csv(report: Dict[str, Any]) -> bytes:
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow(BACKUP_AUDIT_MATRIX_COLUMNS)
    for row in report.get("matrix") or []:
        writer.writerow([_spreadsheet_safe_cell(value) for value in _flatten_target(row)])
    return buffer.getvalue().encode("utf-8")


def export_backup_audit_xlsx(report: Dict[str, Any]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.table import Table, TableStyleInfo

    navy = "17365D"
    blue = "2F75A5"
    pale_blue = "EAF2F8"
    pale_gray = "F4F7FA"
    border_color = "D7E0E8"
    green = "E2F0D9"
    red = "FCE4D6"
    muted = "526477"
    thin_border = Border(bottom=Side(style="thin", color=border_color))
    period = report.get("period") or {}
    summary = report.get("summary") or {}
    availability = report.get("availability") or {}
    size_availability = availability.get("backup_size") or {}
    location_availability = availability.get("storage_location") or {}
    size_summary = summary.get("backup_size") or {}
    size_coverage = size_summary.get("coverage") or {}
    size_metrics = size_availability.get("metrics") or {}

    workbook = Workbook()
    matrix = workbook.active
    matrix.title = "Matriz de respaldos"
    matrix.sheet_view.showGridLines = False
    header_row = 1
    matrix.append(list(BACKUP_AUDIT_XLSX_HEADERS))
    for row in report.get("matrix") or []:
        matrix.append([_spreadsheet_safe_cell(value) for value in _flatten_target(row)])

    last_column = len(BACKUP_AUDIT_MATRIX_COLUMNS)
    last_column_letter = get_column_letter(last_column)
    header_fill = PatternFill("solid", fgColor=navy)
    header_font = Font(name="Aptos", size=9, bold=True, color="FFFFFF")
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for column_index, label in enumerate(BACKUP_AUDIT_XLSX_HEADERS, start=1):
        cell = matrix.cell(row=header_row, column=column_index)
        cell.value = label
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_alignment
        cell.border = Border(bottom=Side(style="medium", color=blue))
    matrix.row_dimensions[header_row].height = 42

    column_widths = (
        20, 28, 22, 24, 15, 17, 14, 11, 18, 17, 22, 24, 12, 13, 12, 13, 14, 13,
        10, 12, 13, 13, 11, 11, 15, 15, 25, 23, 22, 25, 20, 24, 19, 24, 34,
    )
    for column_index, (column_name, width) in enumerate(zip(BACKUP_AUDIT_MATRIX_COLUMNS, column_widths), start=1):
        matrix.column_dimensions[get_column_letter(column_index)].width = width

    for row_number in range(header_row + 1, matrix.max_row + 1):
        fill = PatternFill("solid", fgColor="FFFFFF" if row_number % 2 == 0 else pale_gray)
        for column_index, column_name in enumerate(BACKUP_AUDIT_MATRIX_COLUMNS, start=1):
            cell = matrix.cell(row=row_number, column=column_index)
            cell.fill = fill
            cell.font = Font(name="Aptos", size=9, color="243447")
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = thin_border
            if column_name in ("succeeded_percent", "failed_percent"):
                cell.number_format = '0.00"%"'
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.fill = PatternFill("solid", fgColor=green if column_name == "succeeded_percent" else red)
            elif column_name in ("backup_bytes_processed", "backup_bytes_added"):
                cell.number_format = '#,##0 "bytes"'
                cell.alignment = Alignment(horizontal="right", vertical="center")
            elif column_name in {
                "keep_last", "keep_hourly", "keep_daily", "keep_weekly", "keep_monthly", "keep_yearly",
                "total", "completed", "succeeded", "failed", "other",
                "backup_bytes_processed_reported_jobs", "backup_bytes_processed_unavailable_jobs",
                "backup_bytes_added_reported_jobs", "backup_bytes_added_unavailable_jobs",
            }:
                cell.number_format = "#,##0"
                cell.alignment = Alignment(horizontal="right", vertical="center")
        matrix.row_dimensions[row_number].height = 32

    table_ref = f"A{header_row}:{last_column_letter}{matrix.max_row}"
    if matrix.max_row > header_row:
        audit_table = Table(displayName="BackupAuditMatrix", ref=table_ref)
        audit_table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        matrix.add_table(audit_table)
    else:
        matrix.auto_filter.ref = table_ref
    matrix.freeze_panes = "D2"
    matrix.sheet_properties.pageSetUpPr.fitToPage = True
    matrix.page_setup.orientation = "landscape"
    matrix.page_setup.fitToWidth = 1
    matrix.page_setup.fitToHeight = 0
    matrix.print_title_rows = f"{header_row}:{header_row}"
    matrix.print_area = f"A1:{last_column_letter}{matrix.max_row}"

    executive = workbook.create_sheet("Resumen ejecutivo")
    executive.sheet_view.showGridLines = False
    executive.freeze_panes = "A5"
    for column, width in zip("ABCDEFGH", (22, 18, 18, 18, 18, 18, 18, 18)):
        executive.column_dimensions[column].width = width

    def merge_and_style(sheet, cell_range, value, fill, font, alignment, border=None):
        sheet.merge_cells(cell_range)
        for merged_row in sheet[cell_range]:
            for cell in merged_row:
                cell.fill = fill
                cell.alignment = alignment
                if border is not None:
                    cell.border = border
        anchor = sheet[cell_range.split(":", 1)[0]]
        anchor.value = _spreadsheet_safe_cell(value)
        anchor.font = font
        return anchor

    merge_and_style(
        executive, "A1:H1", "Informe de auditoría de respaldos",
        PatternFill("solid", fgColor=navy), Font(name="Aptos Display", size=22, bold=True, color="FFFFFF"),
        Alignment(vertical="center"),
    )
    executive.row_dimensions[1].height = 40
    merge_and_style(
        executive, "A2:H2",
        f"Período: {period.get('start_date', '')} a {period.get('end_date', '')} ({period.get('timezone', '')})",
        PatternFill("solid", fgColor=pale_blue), Font(name="Aptos", size=11, bold=True, color=navy),
        Alignment(vertical="center"),
    )
    merge_and_style(
        executive, "A3:H3", f"Generado: {report.get('generated_at', '')}",
        PatternFill("solid", fgColor="FFFFFF"), Font(name="Aptos", size=9, color=muted),
        Alignment(vertical="center"),
    )
    executive.row_dimensions[2].height = 26
    executive.row_dimensions[3].height = 21
    executive.row_dimensions[4].height = 10

    cards = (
        (1, 2, "Destinos", summary.get("target_count", 0), pale_blue, navy, "#,##0"),
        (3, 4, "Trabajos de respaldo", summary.get("backup_job_count", 0), pale_blue, navy, "#,##0"),
        (5, 6, "Correctos", summary.get("succeeded_percent"), green, "276221", '0.00"%"'),
        (7, 8, "Fallidos", summary.get("failed_percent"), red, "9C0006", '0.00"%"'),
    )
    for start_column, end_column, label, value, fill, color, number_format in cards:
        start_letter = get_column_letter(start_column)
        end_letter = get_column_letter(end_column)
        card_border = Border(
            left=Side(style="thin", color=border_color), right=Side(style="thin", color=border_color),
            top=Side(style="thin", color=border_color), bottom=Side(style="thin", color=border_color),
        )
        card_fill = PatternFill("solid", fgColor=fill)
        merge_and_style(
            executive, f"{start_letter}5:{end_letter}5", label, card_fill,
            Font(name="Aptos", size=10, bold=True, color=color),
            Alignment(horizontal="center", vertical="center"), card_border,
        )
        value_cell = merge_and_style(
            executive, f"{start_letter}6:{end_letter}7", value if value is not None else "No disponible", card_fill,
            Font(name="Aptos Display", size=22, bold=True, color=color),
            Alignment(horizontal="center", vertical="center"), card_border,
        )
        value_cell.number_format = number_format
    executive.row_dimensions[5].height = 25
    executive.row_dimensions[6].height = 25
    executive.row_dimensions[7].height = 25
    executive.row_dimensions[8].height = 12

    merge_and_style(
        executive, "A9:H9", "Métricas de tamaño y cobertura",
        PatternFill("solid", fgColor=navy), Font(name="Aptos", size=12, bold=True, color="FFFFFF"),
        Alignment(vertical="center"),
    )
    executive.row_dimensions[9].height = 28
    metric_headers = ("Métrica", "Total en bytes", "Cobertura", "Estado")
    for column_index, label in enumerate(metric_headers, start=1):
        cell = executive.cell(row=10, column=column_index, value=label)
        cell.fill = PatternFill("solid", fgColor=blue)
        cell.font = Font(name="Aptos", size=9, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border
    metric_rows = (
        ("bytes_processed", "Bytes de origen procesados"),
        ("bytes_added", "Bytes agregados al backend"),
    )
    for row_number, (metric_key, label) in enumerate(metric_rows, start=11):
        metric = size_metrics.get(metric_key) or {}
        coverage = size_coverage.get(metric_key) or {}
        reported = coverage.get("reported_jobs", metric.get("reported_jobs", 0))
        unavailable = coverage.get("unavailable_jobs", metric.get("unavailable_jobs", 0))
        total = reported + unavailable
        value = size_summary.get(metric_key)
        if not metric.get("available", value is not None) or value is None:
            status = "No disponible"
        elif metric.get("complete", unavailable == 0):
            status = "Completo"
        else:
            status = "Parcial"
        values = (label, value, f"{reported}/{total} trabajos", status)
        for column_index, cell_value in enumerate(values, start=1):
            cell = executive.cell(row=row_number, column=column_index, value=cell_value)
            cell.fill = PatternFill("solid", fgColor="FFFFFF" if row_number % 2 else pale_gray)
            cell.font = Font(name="Aptos", size=10, color="243447")
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            cell.border = thin_border
        executive.cell(row=row_number, column=2).number_format = '#,##0 "bytes"'
        status_cell = executive.cell(row=row_number, column=4)
        status_cell.font = Font(name="Aptos", size=10, bold=True, color="276221" if status == "Completo" else "9C6500")
        status_cell.fill = PatternFill("solid", fgColor=green if status == "Completo" else ("FFF2CC" if status == "Parcial" else pale_gray))
        executive.row_dimensions[row_number].height = 27
    executive.row_dimensions[13].height = 10

    merge_and_style(
        executive, "A14:H14", "Método y calidad de los datos",
        PatternFill("solid", fgColor=navy), Font(name="Aptos", size=12, bold=True, color="FFFFFF"),
        Alignment(vertical="center"),
    )
    size_basis = _spanish_size_basis(bool(size_availability.get("available")))
    location_basis = _spanish_location_basis(bool(location_availability.get("available")))
    merge_and_style(
        executive, "A15:H16", f"Base de las métricas de tamaño: {size_basis}",
        PatternFill("solid", fgColor=pale_gray), Font(name="Aptos", size=9, color=muted),
        Alignment(vertical="top", wrap_text=True), thin_border,
    )
    merge_and_style(
        executive, "A17:H18", f"Base de la ubicación de almacenamiento: {location_basis}",
        PatternFill("solid", fgColor=pale_gray), Font(name="Aptos", size=9, color=muted),
        Alignment(vertical="top", wrap_text=True), thin_border,
    )
    executive.row_dimensions[15].height = 28
    executive.row_dimensions[16].height = 24
    executive.row_dimensions[17].height = 24
    executive.row_dimensions[18].height = 24

    workbook.active = 0
    workbook.properties.title = "Informe de auditoría de respaldos"
    workbook.properties.subject = f"Auditoría de respaldos del {period.get('start_date', '')} al {period.get('end_date', '')}"
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _pdf_escape(text: str) -> bytes:
    escaped = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
    return escaped.encode("latin-1", "replace")


def _pdf_text_width(text: str, font_size: float) -> float:
    narrow = " il.,:;!'|"
    wide = "MW@%&"
    units = sum(0.28 if char in narrow else 0.84 if char in wide else 0.64 if char.isupper() else 0.52 for char in text)
    return units * font_size


def _pdf_wrap_text(text: Any, width: float, font_size: float) -> list[str]:
    normalized = " ".join(str(text if text is not None else "").split())
    if not normalized:
        return [""]
    lines: list[str] = []
    current = ""
    for word in normalized.split(" "):
        candidate = f"{current} {word}".strip()
        if current and _pdf_text_width(candidate, font_size) <= width:
            current = candidate
            continue
        if current:
            lines.append(current)
            current = ""
        fragment = ""
        for character in word:
            candidate = fragment + character
            if fragment and _pdf_text_width(candidate, font_size) > width:
                lines.append(fragment)
                fragment = character
            else:
                fragment = candidate
        current = fragment
    if current:
        lines.append(current)
    return lines or [""]


def _pdf_color(color: Tuple[float, float, float]) -> str:
    return " ".join(f"{component:.3f}" for component in color)


def _pdf_rect(commands: list[bytes], x: float, y: float, width: float, height: float, fill, stroke=None) -> None:
    commands.append(b"q\n")
    if fill is not None:
        commands.append(f"{_pdf_color(fill)} rg\n".encode("ascii"))
    if stroke is not None:
        commands.append(f"{_pdf_color(stroke)} RG 0.6 w\n".encode("ascii"))
    commands.append(f"{x:.2f} {y:.2f} {width:.2f} {height:.2f} re\n".encode("ascii"))
    if fill is not None and stroke is not None:
        commands.append(b"B\n")
    elif fill is not None:
        commands.append(b"f\n")
    elif stroke is not None:
        commands.append(b"S\n")
    commands.append(b"Q\n")


def _pdf_line(commands: list[bytes], x1: float, y1: float, x2: float, y2: float, color) -> None:
    commands.append(b"q\n")
    commands.append(f"{_pdf_color(color)} RG 0.6 w {x1:.2f} {y1:.2f} m {x2:.2f} {y2:.2f} l S\nQ\n".encode("ascii"))


def _pdf_text(
    commands: list[bytes], text: Any, x: float, y: float, font_size: float = 8.0,
    font: str = "F1", color: Tuple[float, float, float] = (0.12, 0.17, 0.23),
) -> None:
    commands.append(
        f"BT /{font} {font_size:.1f} Tf {_pdf_color(color)} rg 1 0 0 1 {x:.2f} {y:.2f} Tm (".encode("ascii")
    )
    commands.append(_pdf_escape(str(text if text is not None else "")))
    commands.append(b") Tj ET\n")


def _pdf_percent(value: Any) -> str:
    if value is None:
        return "No disponible"
    return f"{float(value):.2f}".rstrip("0").rstrip(".") + "%"


def _pdf_human_bytes(value: int) -> str:
    amount = float(value)
    for unit in ("B", "KiB", "MiB", "GiB", "TiB", "PiB"):
        if amount < 1024 or unit == "PiB":
            return f"{amount:,.1f} {unit}" if unit != "B" else f"{int(amount):,} B"
        amount /= 1024
    return f"{int(value):,} B"


def _pdf_metric_details(report: Dict[str, Any], metric_key: str) -> Tuple[str, str, str]:
    summary = report.get("summary") or {}
    size_summary = summary.get("backup_size") or {}
    availability = (report.get("availability") or {}).get("backup_size") or {}
    metric = (availability.get("metrics") or {}).get(metric_key) or {}
    coverage = (size_summary.get("coverage") or {}).get(metric_key) or {}
    reported = coverage.get("reported_jobs", metric.get("reported_jobs", 0))
    unavailable = coverage.get("unavailable_jobs", metric.get("unavailable_jobs", 0))
    total = reported + unavailable
    value = size_summary.get(metric_key)
    if not metric.get("available", value is not None) or value is None:
        status = _spanish_coverage_status("unavailable")
        value_text = "No disponible"
    elif metric.get("complete", unavailable == 0):
        status = _spanish_coverage_status("complete")
        value_text = f"{_pdf_human_bytes(value)} | {value:,} bytes"
    else:
        status = _spanish_coverage_status("partial")
        value_text = f"{_pdf_human_bytes(value)} | {value:,} bytes"
    return value_text, f"{reported}/{total} trabajos", status


def _pdf_draw_card(commands, x, y, width, height, label, value, accent, fill, detail=None, value_size=16.0):
    _pdf_rect(commands, x, y, width, height, fill, (0.84, 0.88, 0.92))
    _pdf_rect(commands, x, y, 4, height, accent)
    _pdf_text(commands, label.upper(), x + 12, y + height - 15, 7.0, "F2", (0.32, 0.39, 0.47))
    _pdf_text(commands, value, x + 12, y + 21, value_size, "F2", accent)
    if detail:
        _pdf_text(commands, detail, x + 12, y + 7, 7.0, "F1", (0.32, 0.39, 0.47))


def _pdf_row_values(row: Dict[str, Any]) -> list[list[str]]:
    server = row.get("server") or {}
    schedule = row.get("schedule") or {}
    retention = row.get("retention") or {}
    counts = row.get("job_counts") or {}
    location = row.get("storage_location") or {}
    backup_size = row.get("backup_size") or {}
    coverage = backup_size.get("coverage") or {}

    def policy_value(key: str) -> str:
        value = retention.get(key)
        if value is None:
            return "n/d"
        if isinstance(value, bool):
            return "Sí" if value else "No"
        return str(value)

    def metric_cell(key: str) -> list[str]:
        metric = coverage.get(key) or {}
        reported = metric.get("reported_jobs", 0)
        unavailable = metric.get("unavailable_jobs", 0)
        total = reported + unavailable
        value = backup_size.get(key)
        if value is None or reported == 0:
            status = _spanish_coverage_status("unavailable")
        elif unavailable:
            status = _spanish_coverage_status("partial")
        else:
            status = _spanish_coverage_status("complete")
        value_text = f"{value:,} bytes" if value is not None else "No disponible"
        return [value_text, f"{status} | {reported}/{total} trabajos"]

    other_statuses = counts.get("other_statuses") or {}
    other_text = ", ".join(
        f"{_spanish_job_status(status)}: {count}" for status, count in sorted(other_statuses.items())
    )
    if not other_text and counts.get("other"):
        other_text = f"Otros: {counts.get('other')}"
    location_lines = [location.get("backend_type") or "No disponible"]
    if location.get("storage_profile_name"):
        location_lines.append(str(location["storage_profile_name"]))
    if location.get("path_storage"):
        location_lines.append(str(location["path_storage"]))
    enabled = row.get("enabled")
    enabled_text = "habilitado" if enabled is True else "deshabilitado" if enabled is False else "estado no disponible"
    return [
        [
            str(row.get("service") or "Servicio sin nombre"),
            f"ID: {row.get('target_id') or 'no disponible'}",
            f"{row.get('runtime_type') or 'no disponible'} | {row.get('backup_strategy') or 'no disponible'}",
            f"Modo: {row.get('backup_mode') or 'no disponible'} | {enabled_text}",
        ],
        [str(server.get("name") or "No disponible"), f"Host: {server.get('host') or 'no disponible'}"],
        [
            f"Programación: {schedule.get('cron_expression') or 'manual'}",
            f"Origen: {_spanish_schedule_source(schedule.get('source'))}",
            f"Zona horaria: {schedule.get('timezone') or 'no disponible'}",
        ],
        [
            str(retention.get("name") or "Sin política"),
            f"Últimos {policy_value('keep_last')} | Cada hora {policy_value('keep_hourly')}",
            f"Cada día {policy_value('keep_daily')} | Cada semana {policy_value('keep_weekly')}",
            f"Cada mes {policy_value('keep_monthly')} | Cada año {policy_value('keep_yearly')}",
            f"Depuración: {policy_value('prune')}",
        ],
        [
            f"{counts.get('completed', 0)}/{counts.get('total', 0)} trabajos completados",
            f"Correctos: {counts.get('succeeded', 0)} | {_pdf_percent(row.get('succeeded_percent'))}",
            f"Fallidos: {counts.get('failed', 0)} | {_pdf_percent(row.get('failed_percent'))}",
            other_text or "Otros: 0",
        ],
        metric_cell("bytes_processed"),
        metric_cell("bytes_added"),
        location_lines,
    ]


def _pdf_layout_row(row: Dict[str, Any], column_widths: Tuple[float, ...]) -> Dict[str, Any]:
    raw_cells = _pdf_row_values(row)
    font_size = 6.8
    cell_lines = [
        [line for raw_line in raw_cell for line in _pdf_wrap_text(raw_line, width - 12, font_size)]
        for raw_cell, width in zip(raw_cells, column_widths)
    ]
    max_lines = max((len(lines) for lines in cell_lines), default=1)
    return {"cells": cell_lines, "height": max(36.0, 12.0 + max_lines * 8.2)}


def _pdf_draw_table_header(commands, x, top, column_widths, labels, colors, height=28.0):
    current_x = x
    for width, label in zip(column_widths, labels):
        _pdf_rect(commands, current_x, top - height, width, height, colors["blue"], colors["white"])
        lines = _pdf_wrap_text(label, width - 10, 7.0)
        line_height = 8.0
        first_y = top - (height - len(lines) * line_height) / 2 - 7.0
        for index, line in enumerate(lines):
            _pdf_text(commands, line, current_x + 5, first_y - index * line_height, 7.0, "F2", colors["white"])
        current_x += width
    return top - height


def _pdf_draw_table_row(commands, top, layout, column_widths, row_index, colors):
    height = layout["height"]
    current_x = 36.0
    stripe = colors["white"] if row_index % 2 == 0 else colors["pale_blue"]
    for column_index, (width, lines) in enumerate(zip(column_widths, layout["cells"])):
        fill = stripe
        if column_index == 4:
            fill = colors["pale_green"] if "Fallidos: 0" in " ".join(lines) else colors["pale_amber"]
        elif column_index in (5, 6):
            joined = " ".join(lines)
            if "No disponible" in joined:
                fill = colors["pale_red"]
            elif "Parcial" in joined:
                fill = colors["pale_amber"]
            elif "Completo" in joined:
                fill = colors["pale_green"]
        _pdf_rect(commands, current_x, top - height, width, height, fill, colors["border"])
        line_y = top - 11.0
        for line_index, line in enumerate(lines):
            font = "F2" if column_index == 0 and line_index == 0 else "F1"
            size = 7.0 if column_index == 0 and line_index == 0 else 6.8
            color = colors["navy"] if font == "F2" else colors["text"]
            _pdf_text(commands, line, current_x + 5, line_y, size, font, color)
            line_y -= 8.2
        current_x += width
    return top - height


def _pdf_content_stream(commands: list[bytes]) -> bytes:
    return b"".join(commands)


def export_backup_audit_pdf(report: Dict[str, Any]) -> bytes:
    page_width, page_height = 842.0, 595.0
    margin = 36.0
    inner_width = page_width - 2 * margin
    column_widths = (102.0, 82.0, 84.0, 94.0, 90.0, 92.0, 92.0, 134.0)
    column_labels = (
        "Servicio / destino", "Servidor", "Programación", "Retención", "Resultados",
        "Bytes de origen procesados", "Bytes agregados al backend", "Destino de almacenamiento",
    )
    colors = {
        "navy": (0.09, 0.21, 0.36),
        "blue": (0.18, 0.40, 0.60),
        "white": (1.0, 1.0, 1.0),
        "text": (0.13, 0.18, 0.24),
        "muted": (0.34, 0.41, 0.49),
        "border": (0.82, 0.86, 0.90),
        "pale_blue": (0.94, 0.97, 0.99),
        "pale_green": (0.91, 0.96, 0.92),
        "pale_amber": (1.0, 0.97, 0.88),
        "pale_red": (0.99, 0.92, 0.91),
        "green": (0.18, 0.44, 0.29),
        "red": (0.67, 0.20, 0.20),
        "amber": (0.64, 0.43, 0.04),
    }
    period = report.get("period") or {}
    summary = report.get("summary") or {}
    availability = report.get("availability") or {}
    size_availability = availability.get("backup_size") or {}
    location_availability = availability.get("storage_location") or {}
    size_summary = summary.get("backup_size") or {}
    matrix = report.get("matrix") or []
    layouts = [_pdf_layout_row(row, column_widths) for row in matrix]

    size_basis = _spanish_size_basis(bool(size_availability.get("available")))
    location_basis = _spanish_location_basis(bool(location_availability.get("available")))
    note_lines = []
    for label, value in (("Base de las métricas", size_basis), ("Ubicación", location_basis)):
        note_lines.extend(_pdf_wrap_text(f"{label}: {value}", inner_width, 7.0))
    first_table_heading_y = 329.0 - len(note_lines) * 9.0 - 8.0
    first_header_top = first_table_heading_y - 12.0
    first_rows_top = first_header_top - 28.0
    first_capacity = first_rows_top - 42.0
    later_rows_top = 464.0
    later_capacity = later_rows_top - 42.0

    pages: list[Dict[str, Any]] = [{"first": True, "rows": []}]
    current_page = pages[0]
    used_height = 0.0
    for layout in layouts:
        capacity = first_capacity if current_page["first"] else later_capacity
        if current_page["rows"] and used_height + layout["height"] > capacity:
            current_page = {"first": False, "rows": []}
            pages.append(current_page)
            used_height = 0.0
            capacity = later_capacity
        current_page["rows"].append(layout)
        used_height += layout["height"]

    def draw_banner(commands, first_page):
        banner_bottom = 500.0 if first_page else 535.0
        _pdf_rect(commands, 0.0, banner_bottom, page_width, page_height - banner_bottom, colors["navy"])
        if first_page:
            _pdf_text(commands, "Informe de auditoría de respaldos", margin, 554, 24.0, "F2", colors["white"])
            _pdf_text(
                commands,
                f"Período: {period.get('start_date', '')} a {period.get('end_date', '')} ({period.get('timezone', '')})",
                margin, 533, 9.0, "F1", colors["white"],
            )
            _pdf_text(commands, f"Generado: {report.get('generated_at', '')}", margin, 516, 7.5, "F1", (0.82, 0.88, 0.94))
        else:
            _pdf_text(commands, "Informe de auditoría de respaldos", margin, 565, 15.0, "F2", colors["white"])
            _pdf_text(
                commands,
                f"{period.get('start_date', '')} a {period.get('end_date', '')} ({period.get('timezone', '')})",
                margin, 546, 8.0, "F1", (0.88, 0.92, 0.96),
            )

    def draw_footer(commands, page_number, page_count):
        _pdf_line(commands, margin, 30.0, page_width - margin, 30.0, colors["border"])
        _pdf_text(commands, "Informe de auditoría de respaldos | Valores exactos y cobertura por métrica", margin, 17.0, 7.0, "F1", colors["muted"])
        _pdf_text(commands, f"Página {page_number} de {page_count}", page_width - margin - 66, 17.0, 7.0, "F2", colors["navy"])

    streams: list[bytes] = []
    for page_index, page in enumerate(pages):
        commands: list[bytes] = []
        first_page = page["first"]
        draw_banner(commands, first_page)
        if first_page:
            gap = 12.0
            card_width = (inner_width - gap * 3) / 4
            kpis = (
                ("Destinos", str(summary.get("target_count", 0)), colors["blue"], colors["pale_blue"]),
                ("Trabajos de respaldo", str(summary.get("backup_job_count", 0)), colors["blue"], colors["pale_blue"]),
                ("Correctos", _pdf_percent(summary.get("succeeded_percent")), colors["green"], colors["pale_green"]),
                ("Fallidos", _pdf_percent(summary.get("failed_percent")), colors["red"], colors["pale_red"]),
            )
            for index, (label, value, accent, fill) in enumerate(kpis):
                _pdf_draw_card(commands, margin + index * (card_width + gap), 427.0, card_width, 56.0, label, value, accent, fill)
            metric_gap = 14.0
            metric_width = (inner_width - metric_gap) / 2
            for index, (metric_key, label, accent) in enumerate((
                ("bytes_processed", "Bytes de origen procesados", colors["blue"]),
                ("bytes_added", "Bytes agregados al backend", colors["green"]),
            )):
                value, coverage, status = _pdf_metric_details(report, metric_key)
                fill = colors["pale_green"] if status == "Completo" else colors["pale_amber"] if status == "Parcial" else colors["pale_red"]
                _pdf_draw_card(
                    commands, margin + index * (metric_width + metric_gap), 352.0, metric_width, 62.0,
                    label, value, accent, fill, f"{status} | {coverage}", 12.0,
                )
            _pdf_text(commands, "Tamaño y calidad de los datos", margin, 337.0, 8.0, "F2", colors["navy"])
            note_y = 325.0
            for line in note_lines:
                _pdf_text(commands, line, margin, note_y, 6.8, "F1", colors["muted"])
                note_y -= 9.0
            _pdf_text(commands, "Matriz de respaldos", margin, first_table_heading_y, 11.0, "F2", colors["navy"])
            row_top = _pdf_draw_table_header(commands, margin, first_header_top, column_widths, column_labels, colors)
        else:
            _pdf_text(commands, "Matriz de respaldos | continuación", margin, 516.0, 10.0, "F2", colors["navy"])
            row_top = _pdf_draw_table_header(commands, margin, 498.0, column_widths, column_labels, colors)
        for row_index, layout in enumerate(page["rows"]):
            row_top = _pdf_draw_table_row(commands, row_top, layout, column_widths, row_index, colors)
        if first_page and not matrix:
            _pdf_text(commands, "No se encontraron destinos de respaldo en este período.", margin + 5, row_top - 18.0, 8.0, "F1", colors["muted"])
        draw_footer(commands, page_index + 1, len(pages))
        streams.append(_pdf_content_stream(commands))

    page_count = len(streams)
    page_obj_nums = [5 + 2 * index for index in range(page_count)]
    content_obj_nums = [6 + 2 * index for index in range(page_count)]
    kids = " ".join(f"{num} 0 R" for num in page_obj_nums)
    body = bytearray(b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n")
    offsets: Dict[int, int] = {}

    def emit(obj_num: int, data: bytes) -> None:
        offsets[obj_num] = len(body)
        body.extend(f"{obj_num} 0 obj\n".encode("ascii"))
        body.extend(data)
        body.extend(b"\nendobj\n")

    emit(1, b"<</Type /Catalog /Pages 2 0 R>>")
    emit(2, f"<</Type /Pages /Kids [{kids}] /Count {page_count}>>".encode("ascii"))
    emit(3, b"<</Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding>>")
    emit(4, b"<</Type /Font /Subtype /Type1 /BaseFont /Helvetica-Bold /Encoding /WinAnsiEncoding>>")
    for index in range(page_count):
        emit(
            page_obj_nums[index],
            (
                f"<</Type /Page /Parent 2 0 R /MediaBox [0 0 842 595] "
                f"/Resources <</Font <</F1 3 0 R /F2 4 0 R>>>> /Contents {content_obj_nums[index]} 0 R>>"
            ).encode("ascii"),
        )
    for index, stream in enumerate(streams):
        emit(
            content_obj_nums[index],
            b"<</Length " + str(len(stream)).encode("ascii") + b">>\nstream\n" + stream + b"\nendstream",
        )

    xref_offset = len(body)
    total_objects = 4 + 2 * page_count
    body.extend(b"xref\n")
    body.extend(f"0 {total_objects + 1}\n".encode("ascii"))
    body.extend(b"0000000000 65535 f \n")
    for obj_num in range(1, total_objects + 1):
        body.extend(f"{offsets[obj_num]:010d} 00000 n \n".encode("ascii"))
    body.extend(b"trailer\n")
    body.extend(f"<</Size {total_objects + 1} /Root 1 0 R>>\n".encode("ascii"))
    body.extend(b"startxref\n")
    body.extend(f"{xref_offset}\n".encode("ascii"))
    body.extend(b"%%EOF")
    return bytes(body)


_EXPORTERS: Dict[str, Callable[[Dict[str, Any]], bytes]] = {
    "json": export_backup_audit_json,
    "csv": export_backup_audit_csv,
    "xlsx": export_backup_audit_xlsx,
    "pdf": export_backup_audit_pdf,
}


def render_backup_audit_export(fmt: str, report: Dict[str, Any]) -> Tuple[str, str, bytes]:
    """Render ``report`` to ``(content_type, filename, body)`` for ``fmt``.

    ``fmt`` must be one of json, csv, xlsx, pdf; any other value raises
    ``ValueError`` so the route maps it to a 400 response.
    """
    normalized = (fmt or "").lower()
    if normalized not in _EXPORTERS:
        raise ValueError(f"unsupported report format: {normalized}")
    period = report.get("period") or {}
    start = str(period.get("start_date") or "")
    end = str(period.get("end_date") or "")
    stamp = f"{start}_{end}" if (start or end) else "report"
    filename = f"backup_audit_{stamp}.{normalized}"
    return _CONTENT_TYPES[normalized], filename, _EXPORTERS[normalized](report)