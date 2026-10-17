import csv
import io
import json
import unittest

from openpyxl import load_workbook

from src.control_plane.application.services.backup_audit_exports import (
    BACKUP_AUDIT_MATRIX_COLUMNS,
    export_backup_audit_csv,
    export_backup_audit_json,
    export_backup_audit_pdf,
    export_backup_audit_xlsx,
    render_backup_audit_export,
)


def _sample_report() -> dict:
    return {
        "report": "backup_audit",
        "generated_at": "2026-03-09T00:00:00Z",
        "period": {
            "start_date": "2026-03-08",
            "end_date": "2026-03-08",
            "timezone": "America/New_York",
        },
        "availability": {
            "backup_size": {
                "available": True,
                "complete": False,
                "basis": (
                    "Restic JSON summary total_bytes_processed (source bytes processed) "
                    "and data_added (backend-reported stored bytes) captured per "
                    "backup.run job; each metric is unavailable independently when a "
                    "job has no exact Restic summary or uses a non-Restic strategy."
                ),
                "metrics": {
                    "bytes_processed": {
                        "source_field": "total_bytes_processed",
                        "label": "source bytes processed",
                        "available": True,
                        "complete": False,
                        "reported_jobs": 1,
                        "unavailable_jobs": 1,
                    },
                    "bytes_added": {
                        "source_field": "data_added",
                        "label": "backend stored bytes",
                        "available": True,
                        "complete": False,
                        "reported_jobs": 1,
                        "unavailable_jobs": 1,
                    },
                },
            },
            "storage_location": {
                "available": True,
                "basis": "target.path_storage and linked storage profile backend_type/name",
            },
        },
        "summary": {
            "target_count": 1,
            "backup_job_count": 2,
            "completed_job_count": 2,
            "succeeded_count": 1,
            "failed_count": 1,
            "other_status_counts": {},
            "unmapped_target_job_count": 0,
            "succeeded_percent": 50.0,
            "failed_percent": 50.0,
            "backup_size": {
                "bytes_processed": 2048,
                "bytes_added": 1024,
                "coverage": {
                    "bytes_processed": {"reported_jobs": 1, "unavailable_jobs": 1},
                    "bytes_added": {"reported_jobs": 1, "unavailable_jobs": 1},
                },
            },
        },
        "matrix": [
            {
                "target_id": "target-a",
                "service": "service-a",
                "server": {"id": "worker-a", "name": "worker-a", "host": "server-a"},
                "runtime_type": "docker",
                "backup_strategy": "restic",
                "backup_mode": "cold",
                "enabled": True,
                "resources": {
                    "compose_project": "project-a",
                    "volumes": ["database:/var/lib/database"],
                    "filesystem_paths": [],
                    "pvc_names": [],
                },
                "schedule": {
                    "cron_expression": "0 2 * * *",
                    "source": "target",
                    "timezone": "America/New_York",
                },
                "retention": {
                    "id": "retention-a",
                    "name": "daily-seven",
                    "keep_last": None,
                    "keep_hourly": None,
                    "keep_daily": 7,
                    "keep_weekly": None,
                    "keep_monthly": None,
                    "keep_yearly": None,
                    "prune": True,
                },
                "job_counts": {
                    "total": 2,
                    "completed": 2,
                    "succeeded": 1,
                    "failed": 1,
                    "other": 0,
                    "other_statuses": {},
                },
                "succeeded_percent": 50.0,
                "failed_percent": 50.0,
                "backup_size": {
                    "bytes_processed": 2048,
                    "bytes_added": 1024,
                    "coverage": {
                        "bytes_processed": {"reported_jobs": 1, "unavailable_jobs": 1},
                        "bytes_added": {"reported_jobs": 1, "unavailable_jobs": 1},
                    },
                },
                "storage_location": {
                    "backend_type": "s3",
                    "storage_profile_name": "prod-s3",
                    "path_storage": "svc/backups",
                },
            }
        ],
    }


class BackupAuditJsonExportTests(unittest.TestCase):
    def test_json_export_roundtrips_the_full_report(self):
        body = export_backup_audit_json(_sample_report())
        self.assertIsInstance(body, bytes)
        self.assertEqual(json.loads(body.decode("utf-8")), _sample_report())


class BackupAuditCsvExportTests(unittest.TestCase):
    def test_csv_has_header_and_one_row_per_target_with_safe_fields_and_size_coverage(self):
        report = _sample_report()
        text = export_backup_audit_csv(report).decode("utf-8")
        rows = list(csv.reader(io.StringIO(text)))
        self.assertEqual(rows[0], list(BACKUP_AUDIT_MATRIX_COLUMNS))
        self.assertEqual(len(rows), 1 + len(report["matrix"]))
        data_row = rows[1]
        self.assertIn("service-a", data_row)
        self.assertIn("server-a", data_row)
        self.assertIn("prod-s3", data_row)
        self.assertIn("svc/backups", data_row)
        # both size metrics and their coverage counts are exposed as columns
        self.assertIn("2048", data_row)
        self.assertIn("1024", data_row)
        columns = dict(zip(BACKUP_AUDIT_MATRIX_COLUMNS, data_row))
        self.assertEqual(columns["backup_bytes_processed"], "2048")
        self.assertEqual(columns["backup_bytes_added"], "1024")
        self.assertEqual(columns["backup_bytes_processed_reported_jobs"], "1")
        self.assertEqual(columns["backup_bytes_processed_unavailable_jobs"], "1")
        self.assertEqual(columns["backup_bytes_added_reported_jobs"], "1")
        self.assertEqual(columns["backup_bytes_added_unavailable_jobs"], "1")


class BackupAuditXlsxExportTests(unittest.TestCase):
    def test_xlsx_is_valid_workbook_with_spanish_headers_and_data_row(self):
        report = _sample_report()
        body = export_backup_audit_xlsx(report)
        self.assertIsInstance(body, bytes)
        self.assertEqual(body[:2], b"PK")  # zip/openxml container
        wb = load_workbook(io.BytesIO(body))
        ws = wb["Matriz de respaldos"]
        self.assertEqual(ws.title, "Matriz de respaldos")

        header_row = None
        for row in ws.iter_rows(values_only=True):
            if row and row[0] == "ID del destino":
                header_row = row
                break
        self.assertIsNotNone(header_row)
        self.assertEqual(
            header_row,
            (
                "ID del destino", "Servicio", "Servidor (worker)", "Host del servidor", "Tipo de ejecución",
                "Estrategia de respaldo", "Modo de respaldo", "Habilitado", "Programación (cron)",
                "Origen de programación", "Zona horaria", "Política de retención", "Últimos",
                "Cada hora", "Cada día", "Cada semana", "Cada mes", "Cada año",
                "Depuración", "Trabajos totales", "Completados", "Correctos", "Fallidos", "Otros",
                "Porcentaje de éxito", "Porcentaje de fallos", "Bytes de origen procesados", "Bytes agregados al backend",
                "Trabajos con bytes procesados", "Trabajos sin bytes procesados", "Trabajos con bytes agregados",
                "Trabajos sin bytes agregados", "Backend de almacenamiento", "Perfil de almacenamiento", "Ruta de almacenamiento",
            ),
        )

        self.assertEqual(ws["A1"].value, "ID del destino")
        self.assertEqual(ws["B1"].value, "Servicio")
        self.assertEqual(ws["C1"].value, "Servidor (worker)")
        self.assertEqual(ws["A2"].value, "target-a")

        data_rows = [row for row in ws.iter_rows(values_only=True) if row and row[0] == "target-a"]
        self.assertEqual(len(data_rows), 1)
        columns = dict(zip(BACKUP_AUDIT_MATRIX_COLUMNS, data_rows[0]))
        self.assertEqual(columns["service"], "service-a")
        self.assertEqual(columns["server_host"], "server-a")
        self.assertEqual(columns["storage_backend_type"], "s3")
        self.assertEqual(columns["storage_profile_name"], "prod-s3")
        self.assertEqual(columns["path_storage"], "svc/backups")
        self.assertEqual(columns["backup_bytes_processed"], 2048)
        self.assertEqual(columns["backup_bytes_added"], 1024)
        self.assertEqual(columns["backup_bytes_processed_reported_jobs"], 1)
        self.assertEqual(columns["backup_bytes_added_unavailable_jobs"], 1)

    def test_xlsx_matrix_has_a_native_filterable_table(self):
        workbook = load_workbook(io.BytesIO(export_backup_audit_xlsx(_sample_report())))
        matrix = workbook["Matriz de respaldos"]
        table = matrix.tables.get("BackupAuditMatrix")
        self.assertIsNotNone(table)
        self.assertEqual(table.ref, f"A1:AI{matrix.max_row}")
        self.assertEqual(table.autoFilter.ref, table.ref)
        self.assertEqual(table.tableStyleInfo.name, "TableStyleMedium2")


class BackupAuditPdfExportTests(unittest.TestCase):
    def test_pdf_is_valid_pdf_containing_spanish_report_text(self):
        report = _sample_report()
        body = export_backup_audit_pdf(report)
        self.assertIsInstance(body, bytes)
        self.assertTrue(body.startswith(b"%PDF-"))
        self.assertTrue(body.rstrip().endswith(b"%%EOF"))
        self.assertIn(b"/Type /Catalog", body)
        self.assertIn(b"/Type /Pages", body)
        self.assertIn(b"(Informe de auditor\xeda de respaldos)", body)
        self.assertIn(b"(Matriz de respaldos)", body)
        self.assertIn(b"(Servicio / destino)", body)
        self.assertIn(b"service-a", body)
        self.assertIn(b"server-a", body)
        # Both byte metrics and their independent coverage remain visible in Spanish.
        self.assertIn(b"2,048 bytes", body)
        self.assertIn(b"1,024 bytes", body)
        self.assertIn(b"1/2 trabajos", body)
        self.assertIn(b"(Parcial | 1/2 trabajos)", body)


class BackupAuditRenderDispatchTests(unittest.TestCase):
    def test_render_returns_content_type_and_dated_filename_for_each_format(self):
        report = _sample_report()
        expected = {
            "json": "application/json",
            "csv": "text/csv; charset=utf-8",
            "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            "pdf": "application/pdf",
        }
        for fmt, content_type in expected.items():
            with self.subTest(fmt=fmt):
                rendered_type, filename, body = render_backup_audit_export(fmt, report)
                self.assertEqual(rendered_type, content_type)
                self.assertTrue(filename.startswith("backup_audit_2026-03-08_2026-03-08."))
                self.assertTrue(filename.endswith(f".{fmt}"))
                self.assertIsInstance(body, bytes)
                self.assertGreater(len(body), 0)

    def test_render_rejects_unsupported_format(self):
        with self.assertRaisesRegex(ValueError, "unsupported report format"):
            render_backup_audit_export("xml", _sample_report())


class BackupAuditSpreadsheetInjectionTests(unittest.TestCase):
    def test_csv_export_neutralizes_formula_like_text(self):
        import csv

        report = _sample_report()
        report["matrix"][0]["service"] = "=1+1"
        rows = list(csv.reader(io.StringIO(export_backup_audit_csv(report).decode("utf-8"))))
        service_index = BACKUP_AUDIT_MATRIX_COLUMNS.index("service")
        self.assertEqual(rows[1][service_index], "'=1+1")

    def test_xlsx_export_stores_formula_like_text_as_a_literal(self):
        report = _sample_report()
        report["matrix"][0]["service"] = "=1+1"
        workbook = load_workbook(io.BytesIO(export_backup_audit_xlsx(report)), data_only=False)
        sheet = workbook["Matriz de respaldos"]
        header_row = next(
            row_number
            for row_number in range(1, sheet.max_row + 1)
            if sheet.cell(row_number, 1).value == "ID del destino"
        )
        service_column = BACKUP_AUDIT_MATRIX_COLUMNS.index("service") + 1
        service_cell = sheet.cell(header_row + 1, service_column)
        self.assertEqual(service_cell.value, "'=1+1")
        self.assertNotEqual(service_cell.data_type, "f")

    def test_xlsx_has_styled_executive_summary_and_readable_matrix(self):
        report = _sample_report()
        wb = load_workbook(io.BytesIO(export_backup_audit_xlsx(report)), data_only=False)
        self.assertEqual(wb.sheetnames, ["Matriz de respaldos", "Resumen ejecutivo"])
        self.assertEqual(wb.active.title, "Matriz de respaldos")

        summary = wb["Resumen ejecutivo"]
        self.assertEqual(summary["A1"].value, "Informe de auditoría de respaldos")
        self.assertTrue(summary["A1"].font.bold)
        self.assertEqual(summary["A1"].fill.fill_type, "solid")
        self.assertIn("A1:H1", {str(ref) for ref in summary.merged_cells.ranges})
        self.assertEqual(summary["A5"].value, "Destinos")
        self.assertEqual(summary["C5"].value, "Trabajos de respaldo")
        self.assertEqual(summary["E5"].value, "Correctos")
        self.assertEqual(summary["G5"].value, "Fallidos")
        self.assertEqual(summary["A6"].value, 1)
        self.assertEqual(summary["C6"].value, 2)
        self.assertEqual(summary["E6"].value, 50.0)
        self.assertEqual(summary["G6"].value, 50.0)
        self.assertEqual(summary["A9"].value, "Métricas de tamaño y cobertura")
        self.assertEqual(summary["A10"].value, "Métrica")
        self.assertEqual(summary["A11"].value, "Bytes de origen procesados")
        self.assertEqual(summary["A12"].value, "Bytes agregados al backend")
        self.assertEqual(summary["C11"].value, "1/2 trabajos")
        self.assertEqual(summary["D11"].value, "Parcial")
        self.assertEqual(summary["B11"].value, 2048)
        self.assertEqual(summary["B12"].value, 1024)
        self.assertEqual(summary["A14"].value, "Método y calidad de los datos")
        self.assertIn("sin estimarlos", summary["A15"].value)
        self.assertIn("perfil de almacenamiento", summary["A17"].value)
        self.assertFalse(summary.sheet_view.showGridLines)

        matrix = wb["Matriz de respaldos"]
        self.assertEqual(matrix.freeze_panes, "D2")
        table = matrix.tables["BackupAuditMatrix"]
        self.assertEqual(table.ref, f"A1:AI{matrix.max_row}")
        self.assertEqual(table.autoFilter.ref, table.ref)
        self.assertEqual(table.tableStyleInfo.name, "TableStyleMedium2")
        self.assertGreaterEqual(matrix.column_dimensions["B"].width, 20)
        self.assertGreater(matrix.row_dimensions[1].height, 20)
        self.assertEqual(matrix["A1"].fill.fill_type, "solid")
        self.assertTrue(matrix["A1"].font.bold)


class BackupAuditPdfLayoutTests(unittest.TestCase):
    def test_pdf_uses_landscape_hierarchy_and_printable_page_footer(self):
        body = export_backup_audit_pdf(_sample_report())
        self.assertIn(b"/MediaBox [0 0 842 595]", body)
        self.assertIn(b"/F2", body)
        self.assertIn(b" rg", body)
        self.assertIn(b"(Matriz de respaldos)", body)
        self.assertIn(b"(Servicio / destino)", body)
        self.assertIn(b"(P\xe1gina 1 de 1)", body)

    def test_pdf_repeats_matrix_header_and_page_number_when_rows_overflow(self):
        import copy

        report = _sample_report()
        template = report["matrix"][0]
        report["matrix"] = []
        for index in range(18):
            row = copy.deepcopy(template)
            row["target_id"] = f"target-{index:02d}"
            row["service"] = f"Long demo service number {index:02d} with a wrapped name"
            report["matrix"].append(row)

        body = export_backup_audit_pdf(report)
        page_count = body.count(b"/Type /Page ")
        self.assertGreater(page_count, 1)
        self.assertEqual(body.count(b"(Servicio / destino)"), page_count)
        self.assertIn(b"(P\xe1gina 1 de", body)
        self.assertIn(b"(P\xe1gina 2 de", body)


if __name__ == "__main__":
    unittest.main()